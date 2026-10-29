"""Bounded spreadsheet parsing, canonical mapping and referential validation."""

import csv
import hashlib
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

TEMPLATES = {
    "clients": ["source_id", "name", "currency"],
    "projects": ["source_id", "name", "client_id", "currency"],
    "contracts": ["source_id", "project_id", "kind", "hourly_rate", "currency"],
    "invoices": [
        "source_id",
        "project_id",
        "client_id",
        "issue_date",
        "due_date",
        "recognition_date",
        "net_amount",
        "tax_amount",
        "discount_amount",
        "status",
        "currency",
    ],
    "invoice_items": ["source_id", "invoice_id", "time_entry_id", "hours", "amount", "coverage_start", "coverage_end"],
    "payments": ["source_id", "payment_date", "amount", "status", "currency"],
    "payment_allocations": ["source_id", "payment_id", "invoice_id", "amount", "currency"],
    "credit_notes": ["source_id", "invoice_id", "amount", "status", "date", "currency"],
    "timesheets": [
        "source_id",
        "project_id",
        "date",
        "hours",
        "billable",
        "billed_hours",
        "write_off",
        "hourly_cost",
        "currency",
    ],
    "expenses": [
        "source_id",
        "project_id",
        "vendor_id",
        "date",
        "amount",
        "category",
        "reference",
        "currency",
        "status",
    ],
    "vendors": ["source_id", "name", "currency"],
}
OPTIONAL = {
    "recognition_date",
    "time_entry_id",
    "coverage_start",
    "coverage_end",
    "project_id:expenses",
    "vendor_id:expenses",
    "hourly_rate",
    "billed_hours",
    "write_off",
    "hourly_cost",
    "tax_amount",
    "discount_amount",
    "hours:invoice_items",
    "reference",
}
REFS = {
    "projects": {"client_id": "clients"},
    "contracts": {"project_id": "projects"},
    "invoices": {"project_id": "projects", "client_id": "clients"},
    "invoice_items": {"invoice_id": "invoices", "time_entry_id": "timesheets"},
    "payment_allocations": {"payment_id": "payments", "invoice_id": "invoices"},
    "credit_notes": {"invoice_id": "invoices"},
    "timesheets": {"project_id": "projects"},
    "expenses": {"project_id": "projects", "vendor_id": "vendors"},
}
MONEY = {"amount", "net_amount", "tax_amount", "discount_amount", "hourly_rate", "hourly_cost"}
NUMERIC = MONEY | {"hours", "billed_hours"}
DATES = {"date", "issue_date", "due_date", "recognition_date", "payment_date", "coverage_start", "coverage_end"}
STATUSES = {
    "invoices": {"issued", "paid", "cancelled", "draft"},
    "payments": {"valid", "reversed", "refund"},
    "credit_notes": {"valid", "cancelled"},
    "expenses": {"valid", "reversed"},
    "contracts": {"hourly", "fixed", "retainer"},
}


def required_fields(kind):
    return [k for k in TEMPLATES[kind] if k not in OPTIONAL and f"{k}:{kind}" not in OPTIONAL]


def canonical_date(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    text = str(value).strip()
    try:
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            raise ValueError("Invalid ISO date")
        return date.fromisoformat(text).isoformat()
    except ValueError:
        raise ValueError("Use an unambiguous ISO date YYYY-MM-DD") from None


def parse_file(name, raw):
    if len(raw) > 8 * 1024 * 1024:
        raise ValueError("Upload limit is 8 MB")
    fingerprint = hashlib.sha256(raw).hexdigest()
    if name.lower().endswith(".csv"):
        try:
            content = raw.decode("utf-8-sig")
            reader = csv.reader(io.StringIO(content), strict=True)
            rows = list(reader)
        except (UnicodeError, csv.Error) as exc:
            raise ValueError("CSV must be UTF-8 with valid quoting") from exc
        sheets = {"CSV": rows}
    elif name.lower().endswith(".xlsx"):
        try:
            with ZipFile(io.BytesIO(raw)) as archive:
                if sum(x.file_size for x in archive.infolist()) > 32 * 1024 * 1024:
                    raise ValueError("Expanded workbook exceeds 32 MB")
                if len(archive.infolist()) > 1000:
                    raise ValueError("Workbook contains too many archive entries")
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False, keep_links=False)
            if len(book.sheetnames) > 20:
                raise ValueError("Maximum 20 workbook sheets")
            sheets = {}
            for sheet in book:
                rows = []
                for row in sheet.iter_rows():
                    if len(rows) > 20000 or len(row) > 60:
                        raise ValueError("Maximum 20,000 rows and 60 columns per sheet")
                    rows.append(
                        [c.value.isoformat()[:10] if isinstance(c.value, (date, datetime)) else c.value for c in row]
                    )
                sheets[sheet.title] = rows
            book.close()
        except (BadZipFile, KeyError, OSError) as exc:
            raise ValueError("Invalid XLSX workbook") from exc
    else:
        raise ValueError("Only .csv and .xlsx files are supported")
    result = {}
    for sheet, rows in sheets.items():
        if not rows:
            continue
        columns = [str(x).strip() if x is not None else "" for x in rows[0]]
        if len(columns) > 60 or len(rows) > 20001:
            raise ValueError("Maximum 20,000 rows and 60 columns")
        if any(not x for x in columns) or len(set(columns)) != len(columns):
            raise ValueError("Column headers must be nonempty and unique")
        if any(len(r) > len(columns) for r in rows[1:]):
            raise ValueError("A row has more cells than the column headers")
        result[sheet] = {
            "columns": columns,
            "rows": [dict(zip(columns, list(r) + [""] * (len(columns) - len(r)))) for r in rows[1:]],
        }
    if not result:
        raise ValueError("No sheets with a header found")
    return fingerprint, result


def validate_rows(kind, sheet, mapping, existing):
    mapping = {k: v for k, v in mapping.items() if v}
    if kind not in TEMPLATES:
        raise ValueError("Unsupported dataset category")
    if any(k not in TEMPLATES[kind] or v not in sheet["columns"] for k, v in mapping.items()):
        raise ValueError("Mapping contains an unknown field or column")
    missing = set(required_fields(kind)) - set(mapping)
    if missing:
        raise ValueError("Map required fields: " + ", ".join(sorted(missing)))
    errors, valid, flagged = [], [], 0
    known = {(k, r["source_id"]): r for k, records in existing.items() for r in records}
    seen = set()
    for number, row in enumerate(sheet["rows"], 2):
        payload = {k: row.get(v, "") for k, v in mapping.items()}
        row_errors = []
        for field in TEMPLATES[kind]:
            value = payload.get(field, "")
            if value is None:
                value = ""
            value = str(value).strip()
            try:
                if value == "" and field in required_fields(kind):
                    raise ValueError("Required value is missing")
                if len(value) > 1000:
                    raise ValueError("Value exceeds 1000 characters")
                if field == "source_id" and len(value) > 160:
                    raise ValueError("Identifier exceeds 160 characters")
                if field in NUMERIC:
                    if field == "hourly_cost" and value == "":
                        payload[field] = None
                        continue
                    value = value or "0"
                    n = Decimal(value)
                    if not n.is_finite() or abs(n) > Decimal("1000000000000"):
                        raise ValueError("Amount must be finite and within 1 trillion")
                    precision = Decimal("0.01")
                    if n != n.quantize(precision):
                        raise ValueError("At most 2 decimal places supported; round explicitly before import")
                    if n < 0 and not (field == "amount" and kind == "payment_allocations"):
                        raise ValueError("Use explicit refund/reversal records instead of negative amounts")
                    value = str(n)
                if field in DATES and value:
                    value = canonical_date(value)
                if field in {"billable", "write_off"}:
                    if value.lower() not in {"", "true", "false", "1", "0", "yes", "no"}:
                        raise ValueError("Expected true or false")
                    value = value.lower() in {"true", "1", "yes"}
                if field == "currency" and value:
                    value = value.upper()
                    if value not in {"USD", "EUR", "GBP", "INR", "CAD", "AUD"}:
                        raise ValueError("Supported currencies: USD, EUR, GBP, INR, CAD, AUD (2-decimal only)")
                if (field == "status" or (kind == "contracts" and field == "kind")) and kind in STATUSES:
                    value = value.lower()
                    if value not in STATUSES[kind]:
                        raise ValueError("Allowed values: " + ", ".join(sorted(STATUSES[kind])))
                payload[field] = value
            except (ValueError, InvalidOperation):
                row_errors.append(
                    {
                        "row": number,
                        "field": field,
                        "message": "Invalid value: " + str(value)[:80]
                        if field in NUMERIC
                        else "Invalid or missing " + field + "; check template and ISO dates",
                    }
                )
        sid = str(payload.get("source_id", ""))
        if sid in seen or (kind, sid) in known:
            row_errors.append(
                {
                    "row": number,
                    "field": "source_id",
                    "message": "Duplicate identifier in this workspace/category or upload",
                }
            )
        seen.add(sid)
        for field, target in REFS.get(kind, {}).items():
            ref = payload.get(field)
            if ref and (target, ref) not in known:
                row_errors.append(
                    {
                        "row": number,
                        "field": field,
                        "message": f"Referenced {target} record does not exist in this workspace",
                    }
                )
            linked = known.get((target, ref), {})
            if linked and payload.get("currency") and linked.get("currency") != payload["currency"]:
                row_errors.append(
                    {
                        "row": number,
                        "field": field,
                        "message": "Cross-currency reference requires FX configuration (not available)",
                    }
                )
        if kind == "invoices" and payload.get("due_date", "") < payload.get("issue_date", ""):
            row_errors.append({"row": number, "field": "due_date", "message": "Due date precedes issue date"})
        if kind == "timesheets":
            try:
                if Decimal(payload.get("billed_hours", "0")) > Decimal(payload.get("hours", "0")):
                    row_errors.append(
                        {"row": number, "field": "billed_hours", "message": "Billed hours exceed logged hours"}
                    )
            except InvalidOperation:
                pass
        if kind == "invoices":
            project = known.get(("projects", payload.get("project_id")), {})
            if project and project.get("client_id") != payload.get("client_id"):
                row_errors.append(
                    {"row": number, "field": "client_id", "message": "Invoice client does not match project client"}
                )
        if kind == "payment_allocations":
            payment = known.get(("payments", payload.get("payment_id")), {})
            try:
                amount = Decimal(payload.get("amount", "0"))
                if (payment.get("status") == "refund" and amount > 0) or (
                    payment.get("status") != "refund" and amount < 0
                ):
                    row_errors.append(
                        {
                            "row": number,
                            "field": "amount",
                            "message": "Refund allocations must be nonpositive; other allocations nonnegative",
                        }
                    )
            except InvalidOperation:
                pass
        if row_errors:
            errors.extend(row_errors)
        else:
            valid.append({"row": number, "payload": payload})
            if kind == "invoices" and not payload.get("recognition_date"):
                flagged += 1
    # Cross-row allocation capacity is checked against the complete accepted batch.
    if kind == "payment_allocations":
        by_payment = {}
        for allocation in existing.get(kind, []) + [x["payload"] for x in valid]:
            pid = allocation["payment_id"]
            by_payment[pid] = by_payment.get(pid, Decimal(0)) + Decimal(allocation["amount"])
        bad = set()
        for pid, allocated in by_payment.items():
            payment = known.get(("payments", pid))
            if payment and (
                (payment["status"] == "refund" and not -Decimal(payment["amount"]) <= allocated <= 0)
                or (payment["status"] != "refund" and not 0 <= allocated <= Decimal(payment["amount"]))
            ):
                bad.add(pid)
        for item in valid[:]:
            if item["payload"]["payment_id"] in bad:
                errors.append(
                    {
                        "row": item["row"],
                        "field": "amount",
                        "message": "Allocation conflicts with payment capacity or refund sign",
                    }
                )
                valid.remove(item)
    return {
        "valid_count": len(valid),
        "rejected_count": len(sheet["rows"]) - len(valid),
        "flagged_count": flagged,
        "errors": errors,
        "preview": [v["payload"] for v in valid[:10]],
        "accepted": valid,
    }
