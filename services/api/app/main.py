"""Local-only MarginProof API. Run Alembic before starting the service."""

import csv
import io
import json
import logging
import os
from contextlib import asynccontextmanager
from datetime import date
from decimal import Decimal
from typing import Annotated, Literal
from xml.sax.saxutils import escape

from fastapi import Depends, FastAPI, File, Header, HTTPException, Query, Request, UploadFile
from fastapi.exceptions import RequestValidationError
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse, Response
from openpyxl import Workbook
from pydantic import BaseModel, ConfigDict, Field, ValidationError
from reportlab.graphics.charts.barcharts import VerticalBarChart
from reportlab.graphics.shapes import Drawing, String
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import delete, func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.analytics import analyze, reconcile_invoice, scenario
from app.database import (
    AnalysisRun,
    Audit,
    Dataset,
    Finding,
    ImportPreview,
    Record,
    RecordLink,
    ScenarioRun,
    SessionLocal,
    Task,
    Workspace,
    get_db,
    now,
    uid,
)
from app.demo import DEMO_ID, seed
from app.imports import REFS, TEMPLATES, parse_file, validate_rows

logger = logging.getLogger("marginproof")
STATES = ["New", "Investigating", "Awaiting Information", "Confirmed", "Resolved", "False Positive", "Archived"]
DEFAULT_START, DEFAULT_END = "2025-11-01", "2026-10-29"


@asynccontextmanager
async def lifespan(app):
    if os.getenv("APP_MODE", "demo") != "demo":
        raise RuntimeError("Production authentication is not implemented. Only APP_MODE=demo is permitted.")
    with SessionLocal() as db:
        seed(db)
        # Completed synchronously; queued/running jobs after interruption are explicitly failed.
        for run in db.scalars(select(AnalysisRun).where(AnalysisRun.status.in_(["queued", "running"]))):
            run.status, run.error, run.finished_at = "failed", "Interrupted before completion; rerun analysis", now()
        db.commit()
        if not db.scalar(select(Finding.id).where(Finding.workspace_id == DEMO_ID).limit(1)):
            run_analysis(db)
    yield


app = FastAPI(
    title="MarginProof AI",
    version="0.1.0",
    description="Development-only fictional demo; not a public multi-tenant service.",
    lifespan=lifespan,
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=False,
    allow_methods=["GET", "POST", "PATCH", "DELETE"],
    allow_headers=["Content-Type", "X-Workspace-Id"],
)


@app.middleware("http")
async def secure_headers(request: Request, call_next):
    request_id = uid()
    if request.method not in {"GET", "HEAD", "OPTIONS"}:
        origin = request.headers.get("origin")
        allowed = {"http://localhost:5173", "http://127.0.0.1:5173", "http://localhost:8000", "http://127.0.0.1:8000"}
        if origin and origin not in allowed:
            return JSONResponse({"error": {"code": 403, "message": "Untrusted origin"}}, status_code=403)
    response = await call_next(request)
    response.headers.update(
        {
            "X-Content-Type-Options": "nosniff",
            "X-Frame-Options": "DENY",
            "Referrer-Policy": "no-referrer",
            "Cache-Control": "no-store",
            "X-Request-Id": request_id,
        }
    )
    logger.info(
        "request method=%s path=%s status=%s request_id=%s",
        request.method,
        request.url.path,
        response.status_code,
        request_id,
    )
    return response


@app.exception_handler(HTTPException)
async def http_error(_, exc):
    return JSONResponse({"error": {"code": exc.status_code, "message": exc.detail}}, status_code=exc.status_code)


@app.exception_handler(ValueError)
async def value_error(_, exc):
    if isinstance(exc, ValidationError):
        return JSONResponse({"error": {"code": 422, "message": "Request validation failed"}}, status_code=422)
    return JSONResponse({"error": {"code": 422, "message": str(exc)}}, status_code=422)


@app.exception_handler(RequestValidationError)
async def validation_error(_, exc):
    return JSONResponse(
        {
            "error": {
                "code": 422,
                "message": "Request validation failed",
                "fields": [{"field": ".".join(str(v) for v in e["loc"]), "message": e["msg"]} for e in exc.errors()],
            }
        },
        status_code=422,
    )


@app.exception_handler(IntegrityError)
async def conflict_error(_, exc):
    return JSONResponse(
        {"error": {"code": 409, "message": "Conflicting record or dependency; refresh and try again"}}, status_code=409
    )


def workspace(x_workspace_id: Annotated[str | None, Header()] = None):
    if x_workspace_id and x_workspace_id != DEMO_ID:
        raise HTTPException(403, "Only the fictional demo workspace is available; private workspace access is disabled")
    return DEMO_ID


DB = Annotated[Session, Depends(get_db)]
WS = Annotated[str, Depends(workspace)]


def audit(db, ws, action, entity, details=None):
    db.add(Audit(workspace_id=ws, action=action, entity_id=entity, details=details or {}))


def get_owned(db, model, identifier, ws):
    item = db.scalar(select(model).where(model.id == identifier, model.workspace_id == ws))
    if not item:
        raise HTTPException(404, "Record not found in this workspace")
    return item


def load_data(db, ws):
    data = {k: [] for k in TEMPLATES}
    for record in db.scalars(select(Record).where(Record.workspace_id == ws)):
        data[record.kind].append(
            {
                **record.payload,
                "id": record.id,
                "dataset_id": record.dataset_id,
                "source_id": record.source_id,
                "source_row": record.source_row,
            }
        )
    return data


def finding_dict(f):
    return {
        **f.payload,
        "id": f.id,
        "status": f.status,
        "owner": f.owner,
        "current": f.current,
        "discovered_at": f.discovered_at,
        "recalculated_at": f.recalculated_at,
    }


def dashboard_data(db, ws, start=DEFAULT_START, end=DEFAULT_END, currency="USD"):
    settings = db.get(Workspace, ws).settings
    result = analyze(load_data(db, ws), start, end, end, currency, settings)
    result["period"] = {"start": start, "end": end}
    fs = {f.key: f for f in db.scalars(select(Finding).where(Finding.workspace_id == ws))}
    # Materialize newly observed period-specific calculations so every displayed
    # finding has a durable evidence/workflow identifier. This caches analysis,
    # never changes source financial records or resets investigation status.
    changed = False
    for payload in result["findings"]:
        if payload["key"] not in fs:
            f = Finding(
                workspace_id=ws,
                key=payload["key"],
                current=False,
                payload={**payload, "analysis_version": "1.0.0", "reporting_period": result["period"]},
            )
            db.add(f)
            db.flush()
            fs[payload["key"]] = f
            changed = True
    if changed:
        db.commit()
    result["findings"] = [
        {
            **payload,
            **(
                {
                    "id": fs[payload["key"]].id,
                    "status": fs[payload["key"]].status,
                    "owner": fs[payload["key"]].owner,
                    "discovered_at": fs[payload["key"]].discovered_at,
                    "recalculated_at": fs[payload["key"]].recalculated_at,
                }
                if payload["key"] in fs
                else {
                    "id": payload["key"],
                    "status": "New",
                    "owner": "Unassigned",
                    "discovered_at": now(),
                    "recalculated_at": now(),
                }
            ),
        }
        for payload in result["findings"]
    ]
    result["metrics"]["unresolved_findings"] = sum(
        f["status"] not in {"Resolved", "False Positive", "Archived"} for f in result["findings"]
    )
    result["metrics"]["completed_investigations"] = sum(
        f["status"] in {"Resolved", "False Positive"} for f in result["findings"]
    )
    result["data_freshness"] = db.scalar(
        select(func.max(Dataset.created_at)).where(Dataset.workspace_id == ws, Dataset.status == "active")
    )
    result["limitations"] = [
        "Fictional local demo; no production authentication.",
        "Recognition uses an explicit date and post-discount invoice net amount; no accrual journal or credit-revenue recognition.",
        "Estimated delivery labor and all external expenses form a contribution estimate, not statutory gross or net profit.",
        "Verified recovered revenue is unavailable: bank-supported verification workflow is not implemented.",
    ]
    result["metrics"]["recovered_revenue"] = None
    return result


def run_analysis(db, ws=DEMO_ID):
    run = AnalysisRun(workspace_id=ws, status="queued")
    db.add(run)
    db.commit()
    run.status = "running"
    db.commit()
    try:
        data = load_data(db, ws)
        dates = [
            r.get(f)
            for records in data.values()
            for r in records
            for f in ["date", "issue_date", "payment_date", "recognition_date"]
            if r.get(f)
        ]
        end = DEFAULT_END  # Fixed reference date makes the fictional demonstration reproducible.
        start = min(dates) if dates else DEFAULT_START
        found = []
        for currency in sorted(
            {r.get("currency") for records in data.values() for r in records if r.get("currency")} or {"USD"}
        ):
            found.extend(analyze(data, start, end, end, currency, db.get(Workspace, ws).settings)["findings"])
        old = {f.key: f for f in db.scalars(select(Finding).where(Finding.workspace_id == ws))}
        for finding in old.values():
            finding.current = False
        for payload in found:
            key = payload["key"]
            payload = {
                **payload,
                "analysis_run_id": run.id,
                "analysis_version": "1.0.0",
                "reporting_period": {"start": start, "end": end},
            }
            if key in old:
                f = old[key]
                # Retain previous evidence in audit history before replacing current calculation.
                audit(
                    db,
                    ws,
                    "finding.recalculated",
                    f.id,
                    {"previous_evidence": f.payload.get("evidence"), "analysis_run_id": run.id},
                )
                f.payload, f.current, f.recalculated_at = payload, True, now()
            else:
                db.add(Finding(workspace_id=ws, key=key, payload=payload))
        run.status, run.finished_at = "completed", now()
        audit(db, ws, "analysis.completed", run.id, {"finding_count": len(found)})
        db.commit()
    except Exception:
        db.rollback()
        run = db.get(AnalysisRun, run.id)
        run.status, run.error, run.finished_at = "failed", "Analysis failed; review data validation and retry", now()
        db.commit()
        raise
    return {"id": run.id, "status": run.status, "finding_count": len(found), "version": run.version}


@app.get("/api/v1/health")
def health(db: DB):
    db.execute(select(1))
    return {"status": "ok", "mode": "development-only", "version": "0.1.0"}


@app.get("/api/v1/workspaces")
def workspaces(db: DB, ws: WS):
    item = db.get(Workspace, ws)
    return [{"id": ws, "name": item.name, "currency": item.settings.get("currency", "USD"), "mode": "Fictional demo"}]


@app.get("/api/v1/dashboard")
def dashboard(
    db: DB,
    ws: WS,
    start: date = date.fromisoformat(DEFAULT_START),
    end: date = date.fromisoformat(DEFAULT_END),
    currency: str = Query("USD", pattern="^(USD|EUR|GBP|INR|CAD|AUD)$"),
):
    return dashboard_data(db, ws, start.isoformat(), end.isoformat(), currency)


@app.get("/api/v1/findings")
def findings(db: DB, ws: WS, status: str | None = None, category: str | None = None):
    items = [
        finding_dict(f)
        for f in db.scalars(select(Finding).where(Finding.workspace_id == ws, Finding.current.is_(True)))
    ]
    return [f for f in items if (not status or f["status"] == status) and (not category or f["category"] == category)]


class FindingUpdate(BaseModel):
    model_config = ConfigDict(extra="forbid")
    status: Literal[
        "New", "Investigating", "Awaiting Information", "Confirmed", "Resolved", "False Positive", "Archived"
    ]
    owner: str | None = Field(None, max_length=120)
    note: str | None = Field(None, max_length=2000)


@app.patch("/api/v1/findings/{identifier}")
def update_finding(identifier: str, body: FindingUpdate, db: DB, ws: WS):
    item = get_owned(db, Finding, identifier, ws)
    previous = item.status
    item.status = body.status
    if body.owner is not None:
        item.owner = body.owner
    audit(
        db, ws, "finding.status", item.id, {"from": previous, "to": body.status, "note": body.note, "owner": item.owner}
    )
    db.commit()
    return finding_dict(item)


@app.get("/api/v1/evidence/{identifier}")
def get_evidence(identifier: str, db: DB, ws: WS):
    f = get_owned(db, Finding, identifier, ws)
    history = db.scalars(
        select(Audit).where(Audit.workspace_id == ws, Audit.entity_id == identifier).order_by(Audit.created_at)
    )
    return {
        "finding": finding_dict(f),
        "history": [{"action": a.action, "details": a.details, "created_at": a.created_at} for a in history],
    }


@app.post("/api/v1/analysis-runs")
def create_analysis(db: DB, ws: WS):
    return run_analysis(db, ws)


@app.get("/api/v1/analysis-runs")
def analysis_runs(db: DB, ws: WS):
    return [
        {"id": r.id, "status": r.status, "created_at": r.created_at, "finished_at": r.finished_at, "error": r.error}
        for r in db.scalars(
            select(AnalysisRun).where(AnalysisRun.workspace_id == ws).order_by(AnalysisRun.created_at.desc()).limit(50)
        )
    ]


def dataset_dict(d):
    return {
        "id": d.id,
        "name": d.name,
        "kind": d.kind,
        "row_count": d.row_count,
        "rejected_count": d.rejected_count,
        "created_at": d.created_at,
        "status": d.status,
    }


@app.get("/api/v1/datasets")
def datasets(db: DB, ws: WS):
    return [
        dataset_dict(d)
        for d in db.scalars(select(Dataset).where(Dataset.workspace_id == ws).order_by(Dataset.created_at.desc()))
    ]


@app.post("/api/v1/imports/preview")
async def preview(db: DB, ws: WS, file: UploadFile = File(...)):
    raw = await file.read(8 * 1024 * 1024 + 1)
    name = (file.filename or "upload.csv").replace("\\", "/").split("/")[-1][:200]
    fingerprint, sheets = parse_file(name, raw)
    if db.scalar(
        select(Dataset.id).where(
            Dataset.workspace_id == ws, Dataset.fingerprint == fingerprint, Dataset.status == "active"
        )
    ):
        raise HTTPException(409, "This file has already been imported. Roll back the existing dataset first.")
    item = ImportPreview(workspace_id=ws, name=name, fingerprint=fingerprint, sheets=sheets)
    db.add(item)
    db.commit()
    first = next(iter(sheets))
    columns = sheets[first]["columns"]
    return {
        "preview_id": item.id,
        "sheets": list(sheets),
        "sheet": first,
        "columns": columns,
        "rows": sheets[first]["rows"][:10],
        "suggested_mapping": {k: k for k in columns if any(k in fields for fields in TEMPLATES.values())},
        "row_count": len(sheets[first]["rows"]),
        "date_format": "ISO YYYY-MM-DD required; ambiguous dates rejected",
    }


class ValidateRequest(BaseModel):
    preview_id: str
    kind: str
    sheet: str | None = None
    mapping: dict[str, str]


@app.get("/api/v1/imports/preview/{identifier}")
def preview_sheet(identifier: str, db: DB, ws: WS, sheet: str):
    item = get_owned(db, ImportPreview, identifier, ws)
    if sheet not in item.sheets:
        raise HTTPException(404, "Sheet not found")
    return {**item.sheets[sheet], "rows": item.sheets[sheet]["rows"][:10]}


@app.post("/api/v1/imports/validate")
def validate(body: ValidateRequest, db: DB, ws: WS):
    item = get_owned(db, ImportPreview, body.preview_id, ws)
    if item.committed:
        raise HTTPException(409, "Preview already committed")
    sheet = body.sheet or next(iter(item.sheets))
    if sheet not in item.sheets:
        raise HTTPException(422, "Sheet not found")
    result = validate_rows(body.kind, item.sheets[sheet], body.mapping, load_data(db, ws))
    item.validation = {**body.model_dump(), "sheet": sheet, "result": result}
    db.commit()
    return {k: v for k, v in result.items() if k != "accepted"} | {"validation_id": item.id}


class CommitRequest(BaseModel):
    validation_id: str


@app.post("/api/v1/imports/commit")
def commit_import(body: CommitRequest, db: DB, ws: WS):
    item = get_owned(db, ImportPreview, body.validation_id, ws)
    if item.committed or not item.validation:
        raise HTTPException(409, "Preview must be validated and may only be committed once")
    if db.scalar(
        select(Dataset.id).where(
            Dataset.workspace_id == ws, Dataset.fingerprint == item.fingerprint, Dataset.status == "active"
        )
    ):
        raise HTTPException(409, "Duplicate import")
    validation = item.validation
    data = load_data(db, ws)
    # Revalidate against current records, never trust a stale preview.
    result = validate_rows(validation["kind"], item.sheets[validation["sheet"]], validation["mapping"], data)
    if not result["valid_count"]:
        raise HTTPException(422, "No valid rows remain; validate again to see errors")
    ds = Dataset(
        workspace_id=ws,
        name=item.name,
        kind=validation["kind"],
        fingerprint=item.fingerprint,
        row_count=result["valid_count"],
        rejected_count=result["rejected_count"],
    )
    db.add(ds)
    db.flush()
    target_records = {(r.kind, r.source_id): r.id for r in db.scalars(select(Record).where(Record.workspace_id == ws))}
    for accepted in result["accepted"]:
        payload = accepted["payload"]
        rec = Record(
            workspace_id=ws,
            dataset_id=ds.id,
            kind=ds.kind,
            source_id=payload["source_id"],
            payload=payload,
            source_row=accepted["row"],
        )
        db.add(rec)
        db.flush()
        for field, target_kind in REFS.get(ds.kind, {}).items():
            if payload.get(field):
                db.add(
                    RecordLink(
                        workspace_id=ws,
                        record_id=rec.id,
                        target_id=target_records[(target_kind, payload[field])],
                        field=field,
                    )
                )
    item.committed = True
    # Full rejected row details remain durable in preview validation and audit summary.
    item.validation = {**validation, "result": result}
    audit(
        db,
        ws,
        "import.committed",
        ds.id,
        {"valid_count": result["valid_count"], "rejected_count": result["rejected_count"], "errors": result["errors"]},
    )
    db.commit()
    run_analysis(db, ws)
    return dataset_dict(ds) | {"errors": result["errors"], "flagged_count": result["flagged_count"]}


@app.delete("/api/v1/datasets/{identifier}")
def rollback(identifier: str, db: DB, ws: WS):
    ds = get_owned(db, Dataset, identifier, ws)
    record_ids = list(db.scalars(select(Record.id).where(Record.workspace_id == ws, Record.dataset_id == identifier)))
    if record_ids:
        dependency = db.scalar(
            select(RecordLink.id)
            .where(
                RecordLink.workspace_id == ws,
                RecordLink.target_id.in_(record_ids),
                RecordLink.record_id.not_in(record_ids),
            )
            .limit(1)
        )
        if dependency:
            raise HTTPException(409, "Other datasets reference these records. Roll back dependent datasets first.")
        db.execute(delete(RecordLink).where(RecordLink.workspace_id == ws, RecordLink.record_id.in_(record_ids)))
        db.execute(delete(Record).where(Record.workspace_id == ws, Record.dataset_id == identifier))
    ds.status = "rolled_back"
    audit(db, ws, "import.rolled_back", identifier, {"records_deleted": len(record_ids)})
    db.commit()
    run_analysis(db, ws)
    return {"status": "rolled_back", "records_deleted": len(record_ids)}


@app.get("/api/v1/templates")
def templates():
    return TEMPLATES


def safe_cell(value):
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=True)
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return "" if value is None else value


def tabular_download(rows, columns, file_format, filename):
    safe_rows = [[safe_cell(r.get(c, "")) for c in columns] for r in rows]
    if file_format == "csv":
        stream = io.StringIO(newline="")
        writer = csv.writer(stream)
        writer.writerow(columns)
        writer.writerows(safe_rows)
        body, media = stream.getvalue().encode("utf-8-sig"), "text/csv"
    elif file_format == "xlsx":
        book = Workbook()
        sheet = book.active
        sheet.title = "MarginProof"
        sheet.append(columns)
        for row in safe_rows:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        from openpyxl.styles import Font, PatternFill

        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="18345A")
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = min(
                65, max(16, max(len(str(c.value or "")) for c in col[:100]) + 2)
            )
        stream = io.BytesIO()
        book.save(stream)
        body, media = stream.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        raise HTTPException(422, "Unsupported export format")
    return Response(
        body, media_type=media, headers={"Content-Disposition": f'attachment; filename="{filename}.{file_format}"'}
    )


@app.get("/api/v1/templates/{kind}")
def template(kind: str, format: Literal["csv", "xlsx"] = "csv"):
    if kind not in TEMPLATES:
        raise HTTPException(404, "Template not found")
    return tabular_download([], TEMPLATES[kind], format, f"marginproof-{kind}-template")


@app.get("/api/v1/records/{kind}")
def records(
    kind: str, db: DB, ws: WS, search: str = "", offset: int = Query(0, ge=0), limit: int = Query(100, ge=1, le=500)
):
    if kind not in TEMPLATES:
        raise HTTPException(404, "Unknown record category")
    items = load_data(db, ws)[kind]
    if search:
        items = [r for r in items if search.lower() in json.dumps(r).lower()]
    return {"items": items[offset : offset + limit], "total": len(items), "offset": offset, "limit": limit}


class ActionBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    title: str = Field(min_length=1, max_length=200)
    owner: str = Field("Unassigned", max_length=120)
    due_date: date | None = None
    status: Literal["Open", "In Progress", "Done"] = "Open"
    note: str = Field("", max_length=4000)
    finding_id: str | None = None
    outcome: str = Field("", max_length=2000)


@app.get("/api/v1/actions")
def actions(db: DB, ws: WS):
    return [
        {"id": a.id, **a.payload, "created_at": a.created_at}
        for a in db.scalars(select(Task).where(Task.workspace_id == ws).order_by(Task.created_at.desc()))
    ]


@app.post("/api/v1/actions")
def create_action(body: ActionBody, db: DB, ws: WS):
    if body.finding_id:
        get_owned(db, Finding, body.finding_id, ws)
    task = Task(workspace_id=ws, payload=body.model_dump(mode="json"))
    db.add(task)
    db.flush()
    audit(db, ws, "action.created", task.id)
    db.commit()
    return {"id": task.id, **task.payload}


@app.patch("/api/v1/actions/{identifier}")
def update_action(identifier: str, body: dict, db: DB, ws: WS):
    item = get_owned(db, Task, identifier, ws)
    validated = ActionBody.model_validate({**item.payload, **body})
    if validated.finding_id:
        get_owned(db, Finding, validated.finding_id, ws)
    item.payload, item.updated_at = validated.model_dump(mode="json"), now()
    audit(db, ws, "action.updated", item.id, {"status": validated.status})
    db.commit()
    return {"id": item.id, **item.payload}


class ScenarioBody(BaseModel):
    rate_change: Decimal = Field(0, ge=-100, le=100)
    cost_change: Decimal = Field(0, ge=-100, le=100)
    unbilled_reduction: Decimal = Field(0, ge=0, le=100)
    delay_days: Decimal = Field(0, ge=0, le=90)


@app.post("/api/v1/scenarios")
def simulate(
    body: ScenarioBody,
    db: DB,
    ws: WS,
    start: date = date.fromisoformat(DEFAULT_START),
    end: date = date.fromisoformat(DEFAULT_END),
    currency: str = "USD",
):
    dash = dashboard_data(db, ws, start.isoformat(), end.isoformat(), currency)
    result = scenario(dash, body.model_dump(mode="json"))
    saved = ScenarioRun(workspace_id=ws, assumptions=body.model_dump(mode="json"), result=result)
    db.add(saved)
    db.commit()
    return result | {"id": saved.id}


class AnalystBody(BaseModel):
    question: str = Field(min_length=3, max_length=1000)


@app.post("/api/v1/analyst")
def analyst(
    body: AnalystBody,
    db: DB,
    ws: WS,
    start: date = date.fromisoformat(DEFAULT_START),
    end: date = date.fromisoformat(DEFAULT_END),
    currency: str = "USD",
):
    d = dashboard_data(db, ws, start.isoformat(), end.isoformat(), currency)
    q = body.question.lower()
    method, rows = "Approved deterministic query", []
    if any(word in q for word in ["overdue", "invoice", "receivable", "payment"]):
        rows = [
            {"id": f["id"], "label": f["title"], "amount": f["amount"]}
            for f in d["findings"]
            if f["category"] in {"overdue", "overdue_invoice", "payment_mismatch", "overpayment"}
        ]
        answer = f"Outstanding receivables are {currency} {d['metrics']['outstanding']}; overdue balance is {currency} {d['metrics']['overdue']}. Balances include partial payments, valid credits and signed refunds."
        method = "Invoice gross payable less credits and signed valid allocations, by contractual due date"
        if "client" in q:
            data = load_data(db, ws)
            names = {r["source_id"]: r["name"] for r in data["clients"]}
            rows, totals = [], {}
            for inv in data["invoices"]:
                if (
                    inv.get("currency") != currency
                    or inv.get("issue_date", "") > end.isoformat()
                    or inv.get("due_date", "") >= end.isoformat()
                ):
                    continue
                balance = reconcile_invoice(
                    inv, data["payments"], data["payment_allocations"], data["credit_notes"], end.isoformat()
                )["outstanding"]
                if Decimal(balance) > 0:
                    name = names.get(inv["client_id"], inv["client_id"])
                    totals[name] = totals.get(name, Decimal(0)) + Decimal(balance)
                    rows.append({"id": inv["id"], "label": f"{name} / {inv['source_id']}", "amount": balance})
            rows.sort(key=lambda r: Decimal(r["amount"]), reverse=True)
            ranked = sorted(totals.items(), key=lambda item: item[1], reverse=True)
            answer += (
                " Highest overdue clients: "
                + (
                    "; ".join(f"{name} ({currency} {amount:.2f})" for name, amount in ranked[:5])
                    if ranked
                    else "none in this scope"
                )
                + ". The source table lists the supporting invoices."
            )
    elif any(word in q for word in ["profit", "margin", "project", "client"]):
        target = Decimal(str(db.get(Workspace, ws).settings.get("target_margin", 35)))
        ranked = sorted(
            [
                r
                for r in d["projects"]
                if r["margin"] is not None and ("target" not in q or Decimal(r["margin"]) < target)
            ],
            key=lambda r: Decimal(r["margin"]),
        )[:8]
        rows = [{"id": r["id"], "label": f"{r['name']} · margin {r['margin']}%", "amount": r["profit"]} for r in ranked]
        answer = f"The contribution estimate is {currency} {d['metrics']['gross_profit']} at {d['metrics']['gross_margin']}% of explicitly recognized revenue. The table ranks the lowest project margins. A causal explanation needs operational context."
        method = "Explicitly recognized invoice net revenue minus external expenses and delivery hours × hourly cost"
    elif any(word in q for word in ["expense", "cost", "spend"]):
        rows = [{"id": r["category"], "label": r["category"], "amount": r["amount"]} for r in d["expense_categories"]]
        answer = f"External expenses total {currency} {d['metrics']['expenses']}; estimated delivery labor is {currency} {d['metrics'].get('delivery_labor')}. Unusual records require review, not automatic rejection."
        method = "Sum of active expense records by date and category; labor separately calculated"
        if len(d["trend"]) >= 2:
            first, last = d["trend"][0], d["trend"][-1]
            change = Decimal(last["expenses"]) - Decimal(first["expenses"])
            answer += f" External expenses changed by {currency} {change:.2f} between {first['month']} and {last['month']}; partial months may not be comparable."
    elif any(word in q for word in ["finding", "investigat", "unbilled", "priorit", "review"]):
        ranked = sorted(d["findings"], key=lambda f: Decimal(f.get("amount") or "0"), reverse=True)[:8]
        rows = [{"id": f["id"], "label": f["title"], "amount": f["amount"]} for f in ranked]
        answer = f"There are {d['metrics']['unresolved_findings']} unresolved findings. Potential hourly unbilled value is {currency} {d['metrics']['unbilled']}, an estimate requiring contract and coverage review. Amounts can overlap and must not be summed as loss."
        method = "Approved finding rules; descending estimable amount (categories may overlap)"
    else:
        answer = "Supported questions cover overdue invoices, project margins, expenses, and investigation priorities. Try: Which projects are below target margin?"
    return {
        "answer": answer,
        "records": rows,
        "method": method,
        "limitations": "Rule-based analyst; no external LLM. Answers use the selected reporting period and explicit stored records. Source identifiers link to tables or findings; no causal claims.",
        "provider": "deterministic",
        "period": d["period"],
    }


@app.get("/api/v1/settings")
def settings(db: DB, ws: WS):
    return db.get(Workspace, ws).settings


class SettingsBody(BaseModel):
    model_config = ConfigDict(extra="forbid")
    business_name: str | None = Field(None, min_length=1, max_length=200)
    currency: Literal["USD", "EUR", "GBP", "INR", "CAD", "AUD"] | None = None
    target_margin: float | None = Field(None, ge=0, le=100)
    overdue_days: int | None = Field(None, ge=0, le=90)
    anomaly_threshold: float | None = Field(None, ge=2, le=10)


@app.patch("/api/v1/settings")
def update_settings(body: SettingsBody, db: DB, ws: WS):
    item = db.get(Workspace, ws)
    item.settings = {**item.settings, **body.model_dump(exclude_none=True)}
    item.name = item.settings["business_name"]
    audit(db, ws, "settings.updated", ws)
    db.commit()
    run_analysis(db, ws)
    return item.settings


REPORTS = {"business", "revenue", "expenses", "profitability", "aging", "findings", "executive", "case"}


@app.get("/api/v1/reports/{report_type}")
def report(
    report_type: str,
    db: DB,
    ws: WS,
    format: Literal["csv", "xlsx", "pdf"] = "pdf",
    start: date = date.fromisoformat(DEFAULT_START),
    end: date = date.fromisoformat(DEFAULT_END),
    currency: str = "USD",
):
    if report_type not in REPORTS:
        raise HTTPException(404, "Unknown report type")
    d = dashboard_data(db, ws, start.isoformat(), end.isoformat(), currency)
    if report_type == "revenue":
        rows = d["clients"]
    elif report_type == "expenses":
        rows = d["expense_categories"]
    elif report_type == "profitability":
        rows = d["projects"]
    elif report_type == "aging":
        rows = d["aging"]
    elif report_type in {"findings", "case"}:
        rows = [
            {
                k: f.get(k)
                for k in ["id", "title", "category", "status", "amount", "currency", "description", "evidence"]
            }
            for f in d["findings"]
        ]
    else:
        rows = [{"metric": k, "value": v} for k, v in d["metrics"].items()]
    generated = now()
    audit(
        db,
        ws,
        "report.exported",
        report_type,
        {"format": format, "start": str(start), "end": str(end), "currency": currency},
    )
    db.commit()
    if format != "pdf":
        metadata = {
            "reporting_period": f"{start} to {end}",
            "currency": currency,
            "generated_at": generated,
            "coverage": d["coverage"],
            "limitations": d["limitations"],
            "source_datasets": [
                r.id for r in db.scalars(select(Dataset).where(Dataset.workspace_id == ws, Dataset.status == "active"))
            ],
        }
        rows = [{**r, **metadata} for r in rows] or [{**metadata, "status": "No data"}]
        return tabular_download(rows, list(rows[0]), format, f"marginproof-{report_type}")
    stream = io.BytesIO()
    doc = SimpleDocTemplate(
        stream, title=f"MarginProof {report_type} report", author="MarginProof AI", leftMargin=40, rightMargin=40
    )
    styles = getSampleStyleSheet()

    def p(text, style="BodyText"):
        return Paragraph(escape(str(text)), styles[style])

    story = [
        p("MARGINPROOF / FINANCIAL INTELLIGENCE", "Heading2"),
        p(f"{report_type.title()} report", "Title"),
        p(db.get(Workspace, ws).name),
        p(f"Fictional local demo | {start} to {end} | {currency}"),
        p(f"Generated {generated}"),
        Spacer(1, 16),
    ]
    display_names = {
        "expenses": "External expenses",
        "gross_profit": "Contribution estimate",
        "gross_margin": "Contribution margin",
        "unbilled": "Potential unbilled work (estimate)",
        "delivery_labor": "Delivery labor (estimate)",
        "recovered_revenue": "Verified recovered revenue",
    }
    summary = [[p("Metric"), p("Selected period / as-of balance")]]
    for key, value in d["metrics"].items():
        if value is None:
            shown = "N/A - verification unavailable" if key == "recovered_revenue" else "N/A - insufficient coverage"
        elif key in {"unresolved_findings", "completed_investigations"}:
            shown = str(value)
        elif key == "gross_margin":
            shown = f"{value}%"
        else:
            shown = f"{currency} {Decimal(str(value)):,.2f}"
        summary.append([p(display_names.get(key, key.replace("_", " ").title())), p(shown)])
    summary_table = Table(summary, colWidths=[240, 275], hAlign="LEFT")
    summary_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EFFA")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
            ]
        )
    )
    story.append(summary_table)
    story += [Spacer(1, 16), p("Revenue and external expense trend", "Heading2")]
    drawing = Drawing(480, 170)
    chart = VerticalBarChart()
    chart.x, chart.y, chart.width, chart.height = 45, 35, 415, 120
    chart.data = [[float(x["revenue"]) for x in d["trend"]], [float(x["expenses"]) for x in d["trend"]]]
    chart.valueAxis.valueMin = 0
    chart.categoryAxis.categoryNames = [x["month"][5:] for x in d["trend"]]
    chart.bars[0].fillColor, chart.bars[1].fillColor = colors.HexColor("#2563EB"), colors.HexColor("#94A3B8")
    if d["trend"]:
        drawing.add(
            String(
                50,
                5,
                f"{currency}: blue = recognized revenue; gray = external expenses",
                fontSize=9,
                fillColor=colors.HexColor("#475569"),
            )
        )
        drawing.add(chart)
        story.append(drawing)
    if report_type not in {"business", "executive"}:
        story += [p("Report detail", "Heading2")]
    for row in rows[:150] if report_type not in {"business", "executive"} else []:
        story.append(p(" | ".join(f"{k}: {v}" for k, v in row.items() if k != "evidence")))
        if row.get("evidence"):
            for ev in row["evidence"]:
                story.append(
                    p(
                        f"Source {ev.get('dataset_id')} / {ev.get('source_id')} / record {ev.get('record_id')}: {ev.get('calculation')} {ev.get('values')}"
                    )
                )
        story.append(Spacer(1, 5))
    if len(rows) > 150:
        story.append(p("PDF detail limited to 150 rows; use CSV/XLSX for full detail."))
    story += [p("Coverage and methodology", "Heading2"), p(json.dumps(d["coverage"]))]
    for m in d["methodology"]:
        story += [p(f"{m['metric']}: {m['formula']}"), p(m["limitations"])]
    story += [
        p("Limitations", "Heading2"),
        *[p(v) for v in d["limitations"]],
        p(
            "Sources: "
            + ", ".join(
                x.id for x in db.scalars(select(Dataset).where(Dataset.workspace_id == ws, Dataset.status == "active"))
            )
        ),
    ]

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#475569"))
        canvas.drawString(40, 25, "MarginProof AI | Fictional demo | Review source evidence before acting")
        canvas.drawRightString(document.pagesize[0] - 40, 25, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return Response(
        stream.getvalue(),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="marginproof-{report_type}.pdf"'},
    )
