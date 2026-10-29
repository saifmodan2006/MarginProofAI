"""Integration tests use an isolated in-memory database, never the local demo DB."""

import base64
import csv
import io
import re
import zlib
from datetime import date
from decimal import Decimal

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import create_engine, event, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app import main
from app.database import Base, Dataset, Finding, ImportPreview, Record, ScenarioRun, Task, Workspace
from app.demo import DEMO_ID, generate

API = "/api/v1"


@pytest.fixture(scope="module")
def harness():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)

    @event.listens_for(engine, "connect")
    def foreign_keys(connection, _):
        connection.execute("PRAGMA foreign_keys=ON")

    Base.metadata.create_all(engine)
    factory = sessionmaker(engine, expire_on_commit=False)

    def db_override():
        with factory() as db:
            yield db

    patch = pytest.MonkeyPatch()
    patch.setattr(main, "SessionLocal", factory)
    patch.setenv("APP_MODE", "demo")
    main.app.dependency_overrides[main.get_db] = db_override
    with TestClient(main.app, raise_server_exceptions=False) as client:
        yield client, factory
    main.app.dependency_overrides.clear()
    patch.undo()
    engine.dispose()


def upload(client, name, content):
    response = client.post(API + "/imports/preview", files={"file": (name, content)})
    assert response.status_code == 200, response.text
    return response.json()


def validate(client, preview, kind, mapping=None, sheet=None):
    return client.post(
        API + "/imports/validate",
        json={
            "preview_id": preview["preview_id"],
            "kind": kind,
            "sheet": sheet,
            "mapping": mapping or {k: k for k in preview["columns"]},
        },
    )


def commit(client, result):
    response = client.post(API + "/imports/commit", json={"validation_id": result.json()["validation_id"]})
    assert response.status_code == 200, response.text
    return response.json()


def csv_bytes(columns, rows):
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(columns)
    writer.writerows(rows)
    return stream.getvalue().encode()


def test_health_security_headers_and_demo_reconciliation(harness):
    client, _ = harness
    response = client.get(API + "/health")
    assert response.status_code == 200
    assert response.json()["status"] == "ok"
    assert response.headers["x-content-type-options"] == "nosniff"
    assert response.headers["cache-control"] == "no-store"
    d = client.get(API + "/dashboard").json()
    raw = generate()
    revenue = sum(Decimal(r["net_amount"]) for r in raw["invoices"])
    expenses = sum(Decimal(r["amount"]) for r in raw["expenses"])
    labor = sum(Decimal(r["hours"]) * Decimal(r["hourly_cost"]) for r in raw["timesheets"])
    # Independent ledger sums, not a second invocation of analyze().
    paid = sum(
        Decimal(r["amount"]) * (-1 if r["status"] == "refund" else 1)
        for r in raw["payments"]
        if r["status"] != "reversed"
    )
    assert Decimal(d["metrics"]["recognized_revenue"]) == revenue
    assert Decimal(d["metrics"]["expenses"]) == expenses
    assert Decimal(d["metrics"]["gross_profit"]) == revenue - expenses - labor
    assert Decimal(d["metrics"]["collected_cash"]) == paid
    assert d["metrics"]["recovered_revenue"] is None
    assert d["coverage"]["currency_conversion"] is False


def test_period_findings_have_durable_workflow_and_evidence_ids(harness):
    client, _ = harness
    result = client.get(API + "/dashboard?start=2026-07-01&end=2026-10-29").json()
    assert result["findings"]
    for finding in result["findings"]:
        identifier = finding["id"]
        assert len(identifier) == 36
        assert client.get(API + f"/evidence/{identifier}").status_code == 200
        response = client.patch(API + f"/findings/{identifier}", json={"status": finding["status"]})
        assert response.status_code == 200


def test_analyst_client_overdue_and_target_queries_cite_actual_records(harness):
    client, factory = harness
    result = client.post(API + "/analyst", json={"question": "Which clients have the highest overdue balances?"}).json()
    assert "Bloom Health" in result["answer"]
    with factory() as db:
        ids = set(db.scalars(select(Record.id).where(Record.workspace_id == DEMO_ID, Record.kind == "invoices")))
    assert result["records"] and all(r["id"] in ids for r in result["records"])
    result = client.post(API + "/analyst", json={"question": "Which projects are below target margin?"}).json()
    assert result["method"] and "no external LLM" not in result["answer"]


def test_csv_custom_mapping_partial_validation_and_dependency_rollback(harness):
    client, _ = harness
    content = csv_bytes(
        ["External key", "Client name", "FX"],
        [["TEST-CLIENT", "=1+1", "USD"], ["TEST-CLIENT", "Duplicate", "USD"], ["BAD-FX", "Bad", "ZZZ"]],
    )
    preview = upload(client, "../clients.csv", content)
    result = validate(
        client, preview, "clients", {"source_id": "External key", "name": "Client name", "currency": "FX"}
    )
    assert result.status_code == 200
    assert result.json()["valid_count"] == 1
    assert result.json()["rejected_count"] == 2
    clients = commit(client, result)
    assert clients["row_count"] == 1 and clients["rejected_count"] == 2
    assert client.post(API + "/imports/preview", files={"file": ("repeat.csv", content)}).status_code == 409
    assert client.post(API + "/imports/commit", json={"validation_id": preview["preview_id"]}).status_code == 409
    project_data = csv_bytes(
        ["source_id", "name", "client_id", "currency"], [["TEST-PROJECT", "Dependent", "TEST-CLIENT", "USD"]]
    )
    projects = commit(client, validate(client, upload(client, "project.csv", project_data), "projects"))
    assert client.delete(API + "/datasets/" + clients["id"]).status_code == 409
    assert client.delete(API + "/datasets/" + projects["id"]).json()["records_deleted"] == 1
    assert client.delete(API + "/datasets/" + clients["id"]).json()["records_deleted"] == 1
    assert client.get(API + "/records/clients", params={"search": "TEST-CLIENT"}).json()["total"] == 0
    assert upload(client, "reimport.csv", content)["row_count"] == 3


def test_xlsx_sheet_selection_dates_and_mapping(harness):
    client, _ = harness
    book = Workbook()
    book.active.title = "Notes"
    book.active.append(["note"])
    book.active.append(["Choose the expenses tab"])
    sheet = book.create_sheet("Expenses")
    columns = ["source_id", "date", "amount", "category", "currency", "status"]
    sheet.append(columns)
    sheet.append(["TEST-XLSX", date(2026, 10, 20), 12.25, "Test", "USD", "valid"])
    sheet.append(["BAD-DATE", "10/11/2026", 20, "Test", "USD", "valid"])
    stream = io.BytesIO()
    book.save(stream)
    preview = upload(client, "multi.xlsx", stream.getvalue())
    assert preview["sheets"] == ["Notes", "Expenses"]
    selected = client.get(API + "/imports/preview/" + preview["preview_id"], params={"sheet": "Expenses"})
    assert selected.json()["rows"][0]["date"] == "2026-10-20"
    result = validate(client, preview, "expenses", {k: k for k in columns}, "Expenses")
    assert result.json()["valid_count"] == 1
    assert result.json()["rejected_count"] == 1
    assert any(e["field"] == "date" for e in result.json()["errors"])
    dataset = commit(client, result)
    assert client.delete(API + "/datasets/" + dataset["id"]).status_code == 200


def test_payment_capacity_refund_sign_and_currency_are_validated(harness):
    client, _ = harness
    columns = ["source_id", "payment_id", "invoice_id", "amount", "currency"]
    rows = [
        ["OVER-CAPACITY", "PAY-001-00-0", "INV-001-00", "1", "USD"],
        ["WRONG-REFUND-SIGN", "REFUND-001", "INV-004-11", "1", "USD"],
        ["WRONG-FX", "PAY-001-00-0", "INV-001-00", "1", "EUR"],
    ]
    result = validate(client, upload(client, "conflicts.csv", csv_bytes(columns, rows)), "payment_allocations")
    assert result.json()["valid_count"] == 0
    assert result.json()["rejected_count"] == 3
    assert (
        client.post(API + "/imports/commit", json={"validation_id": result.json()["validation_id"]}).status_code == 422
    )


def test_refund_payment_and_negative_allocation_import(harness):
    client, _ = harness
    payment = upload(
        client,
        "refund.csv",
        csv_bytes(
            ["source_id", "payment_date", "amount", "status", "currency"],
            [["TEST-REFUND", "2026-10-21", "25", "refund", "USD"]],
        ),
    )
    pay_ds = commit(client, validate(client, payment, "payments"))
    allocation = upload(
        client,
        "refund-allocation.csv",
        csv_bytes(
            ["source_id", "payment_id", "invoice_id", "amount", "currency"],
            [["TEST-REFUND-ALLOC", "TEST-REFUND", "INV-001-11", "-25", "USD"]],
        ),
    )
    alloc_ds = commit(client, validate(client, allocation, "payment_allocations"))
    assert client.delete(API + "/datasets/" + pay_ds["id"]).status_code == 409
    assert client.delete(API + "/datasets/" + alloc_ds["id"]).status_code == 200
    assert client.delete(API + "/datasets/" + pay_ds["id"]).status_code == 200


@pytest.mark.parametrize(
    "method,path,body",
    [
        ("get", "/dashboard", None),
        ("get", "/records/invoices", None),
        ("get", "/datasets", None),
        ("get", "/settings", None),
        ("get", "/findings", None),
        ("get", "/actions", None),
        ("post", "/analysis-runs", {}),
        ("post", "/actions", {"title": "Unauthorized"}),
        ("patch", "/settings", {"business_name": "Unauthorized"}),
    ],
)
def test_private_workspace_header_is_rejected(harness, method, path, body):
    client, _ = harness
    response = client.request(method, API + path, headers={"X-Workspace-Id": "private-tenant"}, json=body)
    assert response.status_code == 403


def test_foreign_objects_cannot_be_read_changed_or_linked(harness):
    client, factory = harness
    with factory() as db:
        db.add(Workspace(id="foreign", name="Private", settings={}))
        db.flush()
        f = Finding(workspace_id="foreign", key="foreign", payload={"title": "Private"})
        task = Task(workspace_id="foreign", payload={"title": "Private"})
        ds = Dataset(workspace_id="foreign", name="Private.csv", kind="clients", fingerprint="private")
        preview = ImportPreview(workspace_id="foreign", name="Private.csv", fingerprint="private", sheets={})
        db.add_all([f, task, ds, preview])
        db.commit()
        ids = f.id, task.id, ds.id, preview.id
    assert client.get(API + "/evidence/" + ids[0]).status_code == 404
    assert client.patch(API + "/findings/" + ids[0], json={"status": "Resolved"}).status_code == 404
    assert client.patch(API + "/actions/" + ids[1], json={"status": "Done"}).status_code == 404
    assert client.delete(API + "/datasets/" + ids[2]).status_code == 404
    assert client.get(API + "/imports/preview/" + ids[3], params={"sheet": "CSV"}).status_code == 404
    assert client.post(API + "/actions", json={"title": "Attempt link", "finding_id": ids[0]}).status_code == 404
    assert all(d["id"] != ids[2] for d in client.get(API + "/datasets").json())


def test_finding_status_owner_evidence_and_action_survive_rerun(harness):
    client, factory = harness
    finding = client.get(API + "/findings").json()[0]
    updated = client.patch(
        API + "/findings/" + finding["id"],
        json={"status": "Investigating", "owner": "Analyst", "note": "Checked bank evidence"},
    )
    assert updated.status_code == 200
    action = client.post(
        API + "/actions", json={"title": "Reconcile bank", "finding_id": finding["id"], "due_date": "2026-11-01"}
    ).json()
    assert (
        client.patch(API + "/actions/" + action["id"], json={"status": "Done", "outcome": "Matched ledger"}).status_code
        == 200
    )
    assert client.post(API + "/analysis-runs").json()["status"] == "completed"
    evidence = client.get(API + "/evidence/" + finding["id"]).json()
    assert evidence["finding"]["status"] == "Investigating"
    assert evidence["finding"]["owner"] == "Analyst"
    assert any(h["action"] == "finding.recalculated" and h["details"]["previous_evidence"] for h in evidence["history"])
    assert any(
        h["action"] == "finding.status" and h["details"]["note"] == "Checked bank evidence" for h in evidence["history"]
    )
    with factory() as db:
        assert db.get(Task, action["id"]).payload["outcome"] == "Matched ledger"


def test_reports_parse_csv_xlsx_pdf_and_escape_spreadsheet_formulas(harness):
    client, factory = harness
    with factory() as db:
        record = db.scalar(select(Record).where(Record.workspace_id == DEMO_ID, Record.kind == "clients"))
        old = record.payload
        record.payload = {**old, "name": '=HYPERLINK("https://example.invalid")'}
        db.commit()
        record_id = record.id
    try:
        response = client.get(API + "/reports/revenue", params={"format": "csv"})
        assert response.status_code == 200
        parsed = list(csv.DictReader(io.StringIO(response.content.decode("utf-8-sig"))))
        assert parsed[0]["name"].startswith("'=HYPERLINK")
        assert parsed[0]["currency"] == "USD"
        assert "reporting_period" in parsed[0] and "limitations" in parsed[0]
        response = client.get(API + "/reports/revenue", params={"format": "xlsx"})
        book = load_workbook(io.BytesIO(response.content))
        sheet = book.active
        header = [cell.value for cell in sheet[1]]
        cell = sheet.cell(2, header.index("name") + 1)
        assert cell.value.startswith("'=HYPERLINK") and cell.data_type != "f"
        assert sheet.freeze_panes == "A2"
        response = client.get(API + "/reports/executive", params={"format": "pdf"})
        assert response.status_code == 200 and response.content.startswith(b"%PDF-")
        assert response.content.rstrip().endswith(b"%%EOF")
        # Decode ReportLab's actual page-content streams, verifying meaningful text.
        text = []
        for stream in re.findall(rb"stream\r?\n(.*?)endstream", response.content, re.S):
            try:
                text.append(zlib.decompress(base64.a85decode(stream.strip(), adobe=True)))
            except (ValueError, zlib.error):
                pass
        assert b"FINANCIAL INTELLIGENCE" in b" ".join(text)
        assert b"Coverage and methodology" in b" ".join(text)
    finally:
        with factory() as db:
            db.get(Record, record_id).payload = old
            db.commit()


def test_scenario_contract_persistence_and_zero_change(harness):
    client, factory = harness
    response = client.post(
        API + "/scenarios", json={"rate_change": 0, "cost_change": 0, "unbilled_reduction": 0, "delay_days": 0}
    )
    assert response.status_code == 200, response.text
    result = response.json()
    for key in ("revenue", "cost", "profit", "margin"):
        assert result["baseline"][key] == result["projected"][key]
    assert result["working_capital_release"] == "0.00"
    assert all(isinstance(s, str) for s in result["assumptions"])
    assert result["sensitivity"] and "profit" in result["sensitivity"][0]
    with factory() as db:
        assert db.get(ScenarioRun, result["id"]).result["projected"] == result["projected"]
    assert client.post(API + "/scenarios", json={"delay_days": 91}).status_code == 422


@pytest.mark.parametrize(
    "name,content",
    [
        ("broken.xlsx", b"not a zip workbook"),
        ("bad.csv", b"\xff\xfe"),
        ("empty.csv", b""),
        ("duplicate.csv", b"id,id\n1,2"),
        ("overflow.csv", b"id\n1,2"),
        ("unsupported.xls", b"legacy"),
    ],
)
def test_malformed_files_have_structured_errors(harness, name, content):
    client, _ = harness
    response = client.post(API + "/imports/preview", files={"file": (name, content)})
    assert response.status_code == 422, response.text
    assert response.json()["error"]["message"]


def test_untrusted_origin_cannot_mutate(harness):
    client, _ = harness
    assert (
        client.post(API + "/actions", json={"title": "blocked"}, headers={"Origin": "https://evil.invalid"}).status_code
        == 403
    )


def test_invalid_action_patch_returns_structured_validation_error(harness):
    client, _ = harness
    action = client.post(API + "/actions", json={"title": "Validation test"}).json()
    response = client.patch(API + "/actions/" + action["id"], json={"status": "Invalid"})
    assert response.status_code == 422, response.text
    assert response.json()["error"]["code"] == 422


def test_unterminated_csv_quote_is_rejected(harness):
    client, _ = harness
    response = client.post(API + "/imports/preview", files={"file": ("quote.csv", b'id,name\n1,"unterminated')})
    assert response.status_code == 422


def test_commit_revalidates_stale_preview_against_current_records(harness):
    client, _ = harness
    columns = ["source_id", "name", "currency"]
    first = upload(client, "first.csv", csv_bytes(columns, [["STALE-PREVIEW", "First", "USD"]]))
    second = upload(client, "second.csv", csv_bytes(columns, [["STALE-PREVIEW", "Second", "USD"]]))
    a = validate(client, first, "clients")
    b = validate(client, second, "clients")
    assert a.json()["valid_count"] == b.json()["valid_count"] == 1
    dataset = commit(client, a)
    response = client.post(API + "/imports/commit", json={"validation_id": b.json()["validation_id"]})
    assert response.status_code == 422
    assert client.get(API + "/records/clients", params={"search": "STALE-PREVIEW"}).json()["total"] == 1
    assert client.delete(API + "/datasets/" + dataset["id"]).status_code == 200


def test_unknown_mapping_and_missing_required_columns_rejected(harness):
    client, _ = harness
    preview = upload(client, "mapping.csv", b"source_id,name,currency\nMAPPING,Name,USD\n")
    for mapping in (
        {"source_id": "source_id"},
        {"source_id": "unknown", "name": "name", "currency": "currency"},
        {"source_id": "source_id", "name": "name", "currency": "currency", "nonexistent": "name"},
    ):
        response = validate(client, preview, "clients", mapping)
        assert response.status_code == 422
        assert response.json()["error"]["message"]


@pytest.mark.parametrize("invalid_date", ["20261020", "2026-10-20junk", "2026-02-30"])
def test_import_dates_require_exact_iso_date(harness, invalid_date):
    client, _ = harness
    columns = ["source_id", "date", "amount", "category", "currency", "status"]
    preview = upload(
        client,
        invalid_date + ".csv",
        csv_bytes(columns, [["EXACT-DATE-" + invalid_date, invalid_date, "10", "Test", "USD", "valid"]]),
    )
    result = validate(client, preview, "expenses")
    assert result.json()["valid_count"] == 0
    assert any(e["field"] == "date" for e in result.json()["errors"])


def test_missing_optional_labor_cost_remains_unknown(harness):
    client, _ = harness
    columns = ["source_id", "project_id", "date", "hours", "billable", "currency"]
    preview = upload(
        client, "unknown-labor.csv", csv_bytes(columns, [["UNKNOWN-LABOR", "PR-001", "2026-10-22", "1", "true", "USD"]])
    )
    result = validate(client, preview, "timesheets")
    assert result.json()["valid_count"] == 1
    assert result.json()["preview"][0]["hourly_cost"] is None
    dataset = commit(client, result)
    try:
        dash = client.get(API + "/dashboard").json()
        assert dash["coverage"]["missing_labor_costs"] == 1
        assert dash["metrics"]["gross_profit"] is None
    finally:
        assert client.delete(API + "/datasets/" + dataset["id"]).status_code == 200
